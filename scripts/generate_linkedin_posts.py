"""Generate LinkedIn post templates for MGVaovao AI translation project."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x0a, 0x66, 0xc2)  # LinkedIn blue
    return p

def post_title(doc, number, angle):
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(f"POST {number}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x0a, 0x66, 0xc2)
    p2 = doc.add_paragraph()
    run2 = p2.add_run(f"Angle : {angle}")
    run2.italic = True
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    doc.add_paragraph()

def post_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p

def note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(f"📌 Note rédaction : {text}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)
    return p

def divider(doc):
    doc.add_paragraph()
    p = doc.add_paragraph("━" * 65)
    p.runs[0].font.color.rgb = RGBColor(0xdd, 0xdd, 0xdd)
    doc.add_paragraph()

# ── COVER PAGE ────────────────────────────────────────────────────────────────
title(doc, "MGVaovao – Maison du Numérique")
title(doc, "Contenus LinkedIn — Projet IA de Traduction Malgache")
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    "6 posts prêts à publier · Storytelling · Annonce · Impact · Vision\n"
    "À adapter selon votre calendrier éditorial"
)
run.italic = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# POST 1 — STORYTELLING / LE PROBLÈME
# ════════════════════════════════════════════════════════════════════════════
post_title(doc, 1, "Storytelling — Le problème qu'on n'ose pas nommer")

post_body(doc, """Il y a une question qu'on nous pose souvent :

« Pourquoi l'IA ne comprend pas le malgache ? »

La réponse est simple — et elle fait mal.

Parce que le malgache ne rapporte pas assez.
Parce que les données d'entraînement n'existent presque pas.
Parce que pendant des années, les langues africaines n'ont pas été prioritaires pour l'industrie technologique mondiale.

Résultat : des millions de Malgaches vivent dans un monde numérique qui ne leur parle pas.
Pas en malgache officiel. Encore moins en betsileo, betsimisaraka ou sakalava.

Chez MGVaovao – Maison du Numérique, nous avons décidé que ce n'était plus acceptable.

Nous avons construit une IA qui écoute, comprend et traduit en malgache — dans ses quatre grands dialectes.

Pas pour remplacer la langue. Pour lui donner sa place dans le monde numérique.

C'est notre combat. Et il commence maintenant.

👇 Suivez notre page pour voir comment on le mène.

#Madagascar #IntelligenceArtificielle #LangageMalgache #Inclusion #Innovation #AfriqueNumérique""")

note(doc, "Idéal pour lancer la présence LinkedIn. Ton engagé, humain, militant.")
divider(doc)

# ════════════════════════════════════════════════════════════════════════════
# POST 2 — ANNONCE PRODUIT
# ════════════════════════════════════════════════════════════════════════════
post_title(doc, 2, "Annonce produit — Ce qu'on a construit")

post_body(doc, """On est fiers d'annoncer quelque chose qu'on travaille depuis longtemps.

🎙️ MGVaovao — le premier système d'IA de traduction vocale pour les langues malgaches.

Voici comment ça fonctionne :

→ Vous parlez en français, en anglais, ou dans une autre langue
→ Notre IA vous écoute en temps réel
→ Elle transcrit, traduit et restitue le résultat en malgache
→ Dans le dialecte de votre choix : officiel, betsileo, betsimisaraka, sakalava

Sous le capot : Whisper pour la reconnaissance vocale, NLLB-200 pour la traduction, et notre propre modèle de synthèse vocale adapté au malgache.

Ce n'est pas juste un outil.
C'est une infrastructure linguistique pour que Madagascar entre dans l'ère de l'IA sans perdre son âme.

La démo est en ligne. La route est longue. Mais on avance.

💬 Des questions sur la technologie ? On répond dans les commentaires.

#MGVaovao #IA #NLP #Madagascar #TechAfrique #TraductionAutomatique #OpenSource""")

note(doc, "Post technique accessible. Bon pour attirer des profils tech et des partenaires institutionnels.")
divider(doc)

# ════════════════════════════════════════════════════════════════════════════
# POST 3 — STORYTELLING HUMAIN / IMPACT
# ════════════════════════════════════════════════════════════════════════════
post_title(doc, 3, "Storytelling humain — Pourquoi ça compte vraiment")

post_body(doc, """Imaginez cette scène.

Une femme dans le Vakinankaratra.
Elle ne parle pas français. Elle ne parle pas anglais.
Elle parle betsileo — la langue de ses parents, de ses enfants, de sa vie.

Elle essaie d'accéder à une information en ligne. Une démarche administrative. Un conseil médical. Une formation.

Chaque outil numérique lui répond dans une langue qui n'est pas la sienne.
Chaque interface la met à distance.
Chaque algorithme l'oublie.

Ce n'est pas une exception. C'est la réalité de millions de Malgaches.

Chez MGVaovao, on développe une IA de traduction pour que cette femme — et tous ceux qui lui ressemblent — ne soient plus exclus du monde numérique.

Parce que l'accès à l'information dans sa propre langue n'est pas un privilège.
C'est un droit.

Notre technologie ne change pas le monde en un jour.
Mais elle pose une brique. Solidement.

Et ça, ça compte.

#Madagascar #InclusionNumérique #Betsileo #LangageIA #TechForGood #Malgache""")

note(doc, "Post émotionnel et humain. Parfait pour toucher un public large, ONG, bailleurs, grand public.")
divider(doc)

# ════════════════════════════════════════════════════════════════════════════
# POST 4 — BEHIND THE SCENES / ÉQUIPE
# ════════════════════════════════════════════════════════════════════════════
post_title(doc, 4, "Behind the scenes — L'équipe derrière le projet")

post_body(doc, """On nous demande souvent : « Qui êtes-vous vraiment ? »

MGVaovao – Maison du Numérique, c'est une équipe malgache.

Pas une startup venue de l'extérieur avec une solution toute faite.
Pas un projet pilote déposé puis oublié.

Des gens du terrain. Des gens qui parlent les langues qu'ils veulent protéger.
Des développeurs, des linguistes, des passionnés — qui travaillent depuis Madagascar, pour Madagascar.

Notre projet d'IA de traduction en langue malgache est né d'une conviction simple :

Les meilleures solutions pour l'Afrique seront faites en Afrique.
Par des Africains. Avec une connaissance intime du contexte local.

On ne prétend pas avoir tout résolu.
On a construit un premier système qui fonctionne.
On continue d'affiner les modèles dialecte par dialecte.
On apprend. On itère. On avance.

Et on cherche des partenaires qui croient comme nous que la langue est le premier levier d'inclusion.

Vous en faites partie ? Écrivez-nous.

#MadeInMadagascar #IA #TraductionMalgache #TeamWork #InnovationAfricaine #Numérique""")

note(doc, "Post d'équipe et de crédibilité. Idéal pour humaniser la marque et attirer des partenaires.")
divider(doc)

# ════════════════════════════════════════════════════════════════════════════
# POST 5 — VISION / LONG TERME
# ════════════════════════════════════════════════════════════════════════════
post_title(doc, 5, "Vision — Où on va dans 5 ans")

post_body(doc, """Dans 5 ans, voilà ce qu'on veut voir :

✅ Un enfant dans le sud de Madagascar qui apprend à lire grâce à une interface vocale en sakalava.

✅ Un médecin en brousse qui reçoit des instructions médicales traduites dans la langue de son patient.

✅ Une femme entrepreneur qui accède à une formation professionnelle en betsileo, depuis son téléphone.

✅ Un touriste qui communique naturellement avec une communauté locale grâce à une traduction instantanée.

Ce n'est pas de la science-fiction.
C'est exactement la direction que prend notre système d'IA de traduction en langue malgache.

Aujourd'hui, on a posé les fondations :
— Reconnaissance vocale adaptée au malgache
— Traduction multi-dialectes (officiel, betsileo, betsimisaraka, sakalava)
— Synthèse vocale en temps réel

Demain, ces briques serviront à construire des applications concrètes dans l'éducation, la santé, l'accès aux services publics.

La langue malgache mérite sa place dans le futur numérique.
On la lui construit.

🔔 Suivez notre page pour ne manquer aucune étape.

#Vision #IA #Madagascar #LangageMalgache #FutureOfAI #TechAfrique #InclusionNumerique""")

note(doc, "Post visionnaire et inspirant. Bon pour les investisseurs, bailleurs et partenaires stratégiques.")
divider(doc)

# ════════════════════════════════════════════════════════════════════════════
# POST 6 — MILESTONE / DÉMO EN LIGNE
# ════════════════════════════════════════════════════════════════════════════
post_title(doc, 6, "Milestone — La démo est en ligne 🎉")

post_body(doc, """C'est une étape qu'on célèbre aujourd'hui.

Notre système d'IA de traduction en temps réel pour les langues malgaches est en ligne.

🎙️ Parlez en français (ou en anglais).
🔄 L'IA transcrit, traduit, et vous répond en malgache — dans le dialecte de votre choix.
🔊 En voix. En temps réel.

Ce moment, on l'a préparé pendant des mois.
Des heures de collecte de données linguistiques.
Des dizaines de cycles d'entraînement de modèles.
Des tests, des erreurs, des corrections.
Et enfin — ça tourne.

Ce n'est pas parfait. On le sait.
Les modèles dialectaux s'affinent encore.
La reconnaissance de certains accents progresse.

Mais c'est réel. C'est fonctionnel. Et c'est fait par une équipe malgache, depuis Madagascar.

On remercie tous nos partenaires qui ont cru en ce projet dès le début.

La suite arrive vite. 🚀

#MGVaovao #Milestone #IA #TraductionMalgache #MadeInMadagascar #Lancement #Innovation""")

note(doc, "Post de lancement produit. À publier le jour où la démo est ouverte au public.")

# ════════════════════════════════════════════════════════════════════════════
# CONSEILS ÉDITORIAUX
# ════════════════════════════════════════════════════════════════════════════
doc.add_page_break()
title(doc, "Conseils pour publier sur LinkedIn")
doc.add_paragraph()

conseils = [
    ("Fréquence recommandée", "1 post par semaine minimum. Commencez par le Post 1 (problème) puis le Post 6 (annonce démo) pour créer une courbe narrative."),
    ("Visuels à prévoir", "Chaque post gagne en portée avec une image ou une courte vidéo (30s de démo de l'interface). Privilégiez des visuels authentiques — pas de stock photos."),
    ("Heure de publication", "Mardi, mercredi ou jeudi entre 8h et 10h du matin (heure de Madagascar) pour toucher l'audience professionnelle locale et internationale."),
    ("Engagement", "Répondez à chaque commentaire dans les 2 premières heures. L'algorithme LinkedIn favorise les posts avec de l'engagement rapide."),
    ("Hashtags", "Limitez à 5-7 hashtags par post. Mélangez des hashtags larges (#IA, #Innovation) et des hashtags de niche (#LangageMalgache, #MGVaovao)."),
    ("Langue", "Ces posts sont en français. Envisagez une version en anglais pour toucher l'audience africaine anglophone et les partenaires internationaux."),
]

for titre_conseil, texte_conseil in conseils:
    p = doc.add_paragraph()
    run1 = p.add_run(f"▸ {titre_conseil} : ")
    run1.bold = True
    run1.font.size = Pt(11)
    run2 = p.add_run(texte_conseil)
    run2.font.size = Pt(11)

# ── SAVE ─────────────────────────────────────────────────────────────────────
output = "C:/Github Repositories/mgvaovao/LINGUA_LinkedIn_Posts_MGVaovao.docx"
doc.save(output)
print(f"Saved: {output}")

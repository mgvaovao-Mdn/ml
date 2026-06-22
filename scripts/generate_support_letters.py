"""Generate LINGUA AFRICA MASHAKANE support letter templates as a Word document."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

def heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x1a, 0x56, 0x9e)
    return p

def subheading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return p

def body(doc, text, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.italic = italic
    return p

def field(doc, text):
    """Placeholder field shown in grey italics."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.italic = True
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    return p

def separator(doc):
    doc.add_paragraph()
    p = doc.add_paragraph("─" * 70)
    p.runs[0].font.color.rgb = RGBColor(0xcc, 0xcc, 0xcc)
    doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
# TITRE GÉNÉRAL
# ════════════════════════════════════════════════════════════════════════════
heading(doc, "MODÈLES DE LETTRES DE SOUTIEN")
heading(doc, "MGVaovao – Maison du Numérique")
heading(doc, "Candidature LINGUA AFRICA MASHAKANE")
doc.add_paragraph()
body(doc,
    "Ces modèles sont fournis à titre de brainstorming. Chaque partenaire est invité "
    "à les adapter, les imprimer sur papier à en-tête et les signer avant envoi.",
    italic=True)
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# LETTRE 1 — YANKY
# ════════════════════════════════════════════════════════════════════════════
heading(doc, "LETTRE 1 — YANKY")
subheading(doc, "Angle : Innovation technologique & écosystème numérique")
doc.add_paragraph()

field(doc, "YANKY\n[Adresse]\n[Ville, Madagascar]\n[Date]")
doc.add_paragraph()

body(doc, "Objet : Lettre de soutien à la candidature de MGVaovao - Maison du Numérique "
     "au concours LINGUA AFRICA MASHAKANE")
doc.add_paragraph()

body(doc, "Madame, Monsieur les membres du jury,")
doc.add_paragraph()

body(doc,
    "En tant que partenaire actif de MGVaovao - Maison du Numérique, nous avons le plaisir "
    "d'apporter notre soutien plein et entier à leur candidature au concours LINGUA AFRICA MASHAKANE.")

body(doc,
    "Nous avons pu observer de près le travail remarquable accompli par l'équipe de MGVaovao "
    "dans le développement d'une solution d'intelligence artificielle dédiée aux langues malgaches. "
    "Dans un contexte où les outils numériques sont quasi exclusivement conçus pour des langues "
    "internationales, MGVaovao relève le défi de rendre la technologie accessible dans les quatre "
    "grandes variantes dialectales du malgache : le malgache officiel, le betsileo, "
    "le betsimisaraka et le sakalava.")

body(doc,
    "Ce projet incarne exactement ce que nous défendons au sein de notre écosystème : "
    "l'innovation au service des réalités locales, portée par des Malgaches pour les Malgaches.")

body(doc,
    "Nous sommes convaincus que MGVaovao représente la Madagascar numérique de demain et que "
    "ce concours constitue une opportunité unique d'accélérer leur impact à l'échelle du "
    "continent africain.")

body(doc,
    "Nous soutenons sans réserve cette candidature et restons disponibles pour toute "
    "information complémentaire.")
doc.add_paragraph()

body(doc, "Veuillez agréer, Madame, Monsieur, l'expression de nos sincères salutations.")
doc.add_paragraph()
field(doc, "[Nom & Signature]\nPour YANKY")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# LETTRE 2 — YAS / Ampela Online
# ════════════════════════════════════════════════════════════════════════════
heading(doc, "LETTRE 2 — YAS / Ampela Online")
subheading(doc, "Angle : Autonomisation des femmes & inclusion numérique")
doc.add_paragraph()

field(doc, "YAS / Ampela Online\n[Adresse]\n[Ville, Madagascar]\n[Date]")
doc.add_paragraph()

body(doc, "Objet : Lettre de soutien à la candidature de MGVaovao au concours LINGUA AFRICA MASHAKANE")
doc.add_paragraph()

body(doc, "Madame, Monsieur les membres du jury,")
doc.add_paragraph()

body(doc,
    "La plateforme YAS / Ampela Online, engagée dans l'autonomisation des femmes et des jeunes "
    "par le numérique, soutient avec fierté la candidature de MGVaovao - Maison du Numérique "
    "au concours LINGUA AFRICA MASHAKANE.")

body(doc,
    "L'une des barrières majeures à la participation des femmes malgaches dans l'espace numérique "
    "est la langue. La majorité des ressources, formations et outils disponibles en ligne sont en "
    "français ou en anglais, excluant de facto des millions de femmes qui s'expriment principalement "
    "dans leur dialecte maternel.")

body(doc,
    "Le projet MGVaovao s'attaque directement à cette barrière en développant un outil de traduction "
    "et de synthèse vocale en temps réel, couvrant plusieurs dialectes malgaches. Cela ouvre "
    "concrètement la porte à une inclusion numérique qui ne laisse personne au bord du chemin.")

body(doc,
    "Nous avons collaboré avec l'équipe de MGVaovao et sommes témoins de leur sérieux, de leur "
    "vision à long terme et de leur ancrage profond dans les réalités du terrain malgache.")

body(doc,
    "Le concours LINGUA AFRICA MASHAKANE est une scène idéale pour faire rayonner ce projet au "
    "niveau africain et nous soutenons pleinement cette candidature.")
doc.add_paragraph()

body(doc, "Avec nos meilleures salutations,")
doc.add_paragraph()
field(doc, "[Nom & Signature]\nPour YAS / Ampela Online")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# LETTRE 3 — EJERY
# ════════════════════════════════════════════════════════════════════════════
heading(doc, "LETTRE 3 — EJERY")
subheading(doc, "Angle : Culture, patrimoine linguistique & jeunesse")
doc.add_paragraph()

field(doc, "EJERY\n[Adresse]\n[Ville, Madagascar]\n[Date]")
doc.add_paragraph()

body(doc, "Objet : Lettre de soutien à MGVaovao - Maison du Numérique pour le concours LINGUA AFRICA MASHAKANE")
doc.add_paragraph()

body(doc, "Madame, Monsieur les membres du jury,")
doc.add_paragraph()

body(doc,
    "EJERY apporte son soutien enthousiaste à la candidature de MGVaovao - Maison du Numérique "
    "au concours LINGUA AFRICA MASHAKANE.")

body(doc,
    "La langue malgache est un patrimoine vivant, riche de ses variantes régionales qui portent "
    "chacune l'identité, la culture et la mémoire d'une communauté. Aujourd'hui, ce patrimoine "
    "est menacé par une numérisation qui l'ignore. Les jeunes générations évoluent dans un monde "
    "numérique qui ne parle pas leur langue.")

body(doc,
    "MGVaovao change cela. Leur système d'intelligence artificielle — capable d'écouter, de "
    "comprendre et de traduire en malgache dans ses différents dialectes — est un acte de "
    "valorisation culturelle autant qu'une innovation technologique. C'est redonner au malgache "
    "sa place dans l'ère numérique.")

body(doc,
    "Nous croyons profondément que la jeunesse malgache mérite des outils numériques qui lui "
    "ressemblent et parlent sa langue. MGVaovao construit ces outils avec rigueur et ambition.")

body(doc,
    "Nous encourageons vivement le jury à soutenir ce projet porteur de sens pour Madagascar "
    "et pour l'Afrique.")
doc.add_paragraph()

body(doc, "Sincèrement,")
doc.add_paragraph()
field(doc, "[Nom & Signature]\nPour EJERY")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# LETTRE 4 — DOWN SYNDROME MADAGASCAR
# ════════════════════════════════════════════════════════════════════════════
heading(doc, "LETTRE 4 — DOWN SYNDROME MADAGASCAR")
subheading(doc, "Angle : Accessibilité, inclusion & technologie pour tous")
doc.add_paragraph()

field(doc, "DOWN SYNDROME MADAGASCAR\n[Adresse]\n[Ville, Madagascar]\n[Date]")
doc.add_paragraph()

body(doc, "Objet : Lettre de soutien à la candidature de MGVaovao au concours LINGUA AFRICA MASHAKANE")
doc.add_paragraph()

body(doc, "Madame, Monsieur les membres du jury,")
doc.add_paragraph()

body(doc,
    "L'association Down Syndrome Madagascar soutient avec conviction la candidature de "
    "MGVaovao - Maison du Numérique au concours LINGUA AFRICA MASHAKANE.")

body(doc,
    "Notre association œuvre quotidiennement pour l'inclusion des personnes en situation de "
    "handicap dans la société malgache. L'une des réalités que nous affrontons est celle de la "
    "communication : pour de nombreuses personnes que nous accompagnons, le français est une "
    "barrière supplémentaire qui s'ajoute aux difficultés existantes. Avoir accès à l'information, "
    "aux soins, à l'éducation dans sa propre langue maternelle n'est pas un luxe — c'est un "
    "droit fondamental.")

body(doc,
    "La technologie développée par MGVaovao — permettant la transcription, la traduction et la "
    "synthèse vocale en malgache — représente une avancée concrète pour l'accessibilité. Elle "
    "ouvre la possibilité de créer des interfaces vocales, des outils de communication augmentée, "
    "et des ressources pédagogiques dans les langues que comprennent réellement les personnes "
    "que nous soutenons.")

body(doc,
    "MGVaovao ne développe pas simplement un outil technologique. Ils construisent un pont entre "
    "le numérique et ceux qui en sont aujourd'hui exclus. C'est une mission que nous partageons "
    "pleinement.")

body(doc,
    "Nous apportons notre soutien sans réserve à cette candidature et espérons qu'elle recevra "
    "la reconnaissance qu'elle mérite.")
doc.add_paragraph()

body(doc, "Chaleureusement,")
doc.add_paragraph()
field(doc, "[Nom & Signature]\nPour Down Syndrome Madagascar")

# ════════════════════════════════════════════════════════════════════════════
# SAVE
# ════════════════════════════════════════════════════════════════════════════
output = "C:/Github Repositories/mgvaovao/LINGUA_Lettres_Soutien_Partenaires.docx"
doc.save(output)
print(f"Saved: {output}")
